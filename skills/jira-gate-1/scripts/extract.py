#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""附件文本抽取脚本（Jira 需求准入预检用）

用途：把 Jira 需求单上的附件转成纯文本，供准入检查逐项阅读。
运行：用任意带以下依赖的 Python 3 执行（Codex 内置运行时已包含全部依赖）：
    python extract.py <文件或目录> [--max-chars 20000] [--out 输出文件] [--unpack-dir 解压目录]
依赖：charset-normalizer、python-docx、pdfplumber、pypdf、openpyxl、python-pptx

支持格式：
    .txt/.md/.markdown/.json/.xml/.html/.htm/.csv/.log/.yaml/.yml/.sql/.properties  → 直接按编码读取
    .docx  → python-docx（含表格）
    .pdf   → pdfplumber（失败时退回 pypdf）
    .xlsx/.xlsm → openpyxl
    .pptx  → python-pptx
    .zip / .rar / .7z / .tar / .tar.gz 等 → 先解压再递归解析（rar 走系统自带 tar.exe，实测支持 RAR5）
    图片   → 提示改用图像查看能力，不做 OCR
不支持的格式会明确提示“无法解析”，避免把解析失败当成“附件没有内容”。
"""

import argparse
import bz2
import gzip
import lzma
import os
import shutil
import subprocess
import sys
import tarfile
import zipfile

TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".json", ".xml", ".html", ".htm", ".csv",
    ".log", ".yaml", ".yml", ".sql", ".properties", ".ini", ".conf", ".tsv",
}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff"}
UNSUPPORTED_HINT = {".doc": "旧版 .doc 无法直接解析，请让产品经理另存为 .docx 或 .pdf"}

# 归档（压缩包）相关配置
ARCHIVE_DIR_NAME = "__unpacked"          # 默认解压目录名，遍历时会跳过，避免重复解析
MAX_ARCHIVE_DEPTH = 3                    # 压缩包嵌套层数上限
TAR_LIKE_SUFFIXES = (".tar.gz", ".tar.bz2", ".tar.xz", ".tgz", ".tbz2", ".txz", ".tar")
SINGLE_COMPRESS_MODULES = {".gz": gzip, ".bz2": bz2, ".xz": lzma}
STRIP_SUFFIXES = TAR_LIKE_SUFFIXES + (".zip", ".rar", ".7z", ".gz", ".bz2", ".xz")

# rar / 7z 没有可用的纯 Python 解压实现，按顺序尝试系统命令。
# Windows 10+ 自带的 bsdtar（tar.exe）实测能解 RAR5，是没装 7-Zip / WinRAR 时的主力方案。
EXTERNAL_UNPACKERS = {
    "rar": (
        ("tar", ["-xf", "{archive}", "-C", "{dest}"]),
        ("7z", ["x", "-y", "-o{dest}", "{archive}"]),
        ("7za", ["x", "-y", "-o{dest}", "{archive}"]),
        ("unrar", ["x", "-o+", "{archive}", "{dest}"]),
    ),
    "7z": (
        ("7z", ["x", "-y", "-o{dest}", "{archive}"]),
        ("7za", ["x", "-y", "-o{dest}", "{archive}"]),
        ("tar", ["-xf", "{archive}", "-C", "{dest}"]),
    ),
}


def read_text_file(path):
    """按探测到的编码读取纯文本文件。"""
    try:
        from charset_normalizer import from_path
        best = from_path(path).best()
        if best is not None:
            return str(best)
    except Exception:
        pass
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
        try:
            with open(path, "r", encoding=encoding) as fh:
                return fh.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    with open(path, "rb") as fh:
        return fh.read().decode("utf-8", errors="replace")


def read_docx(path):
    """提取 docx 的段落与表格内容。"""
    from docx import Document
    doc = Document(path)
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
    for index, table in enumerate(doc.tables, 1):
        lines.append(f"[表格 {index}]")
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def read_pdf(path):
    """提取 pdf 文本，逐页拼接。"""
    try:
        import pdfplumber
        lines = []
        with pdfplumber.open(path) as pdf:
            for index, page in enumerate(pdf.pages, 1):
                text = (page.extract_text() or "").strip()
                lines.append(f"[第 {index} 页]")
                lines.append(text)
        joined = "\n".join(lines).strip()
        if joined:
            return joined
    except Exception as exc:  # noqa: BLE001
        fallback_reason = str(exc)
    else:
        fallback_reason = "pdfplumber 未提取到文本（可能是扫描件）"
    try:
        from pypdf import PdfReader
        reader = PdfReader(path)
        text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
        if text:
            return text
    except Exception as exc:  # noqa: BLE001
        return f"[PDF 解析失败] {fallback_reason}；pypdf 也失败：{exc}"
    return f"[PDF 解析失败] {fallback_reason}"


def read_xlsx(path):
    """提取 Excel 每个工作表的单元格内容。"""
    from openpyxl import load_workbook
    workbook = load_workbook(path, data_only=True, read_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"[工作表: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if cell is None else str(cell).strip() for cell in row]
            if any(cells):
                lines.append(" | ".join(cells))
    workbook.close()
    return "\n".join(lines)


def read_pptx(path):
    """提取 pptx 每页文本框内容。"""
    from pptx import Presentation
    presentation = Presentation(path)
    lines = []
    for index, slide in enumerate(presentation.slides, 1):
        lines.append(f"[第 {index} 页]")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        lines.append(text)
    return "\n".join(lines)


def archive_kind(path):
    """识别归档类型，返回 'tar' / 'zip' / 'rar' / '7z' / 'single'，非归档返回 None。"""
    name = os.path.basename(path).lower()
    if name.endswith(TAR_LIKE_SUFFIXES):
        return "tar"
    suffix = os.path.splitext(name)[1]
    if suffix == ".zip":
        return "zip"
    if suffix == ".rar":
        return "rar"
    if suffix == ".7z":
        return "7z"
    if suffix in SINGLE_COMPRESS_MODULES:
        return "single"
    return None


def _is_within(dest_dir, target):
    """判断解压目标路径是否落在解压目录内，防止压缩包路径穿越（zip slip）。"""
    dest_dir = os.path.abspath(dest_dir)
    target = os.path.abspath(target)
    return target == dest_dir or target.startswith(dest_dir + os.sep)


def _zip_member_name(info):
    """zip 中文名兼容：未置 UTF-8 标志位时按 cp437 解码再按 GBK 还原。"""
    if info.flag_bits & 0x800:
        return info.filename
    try:
        return info.filename.encode("cp437").decode("gbk")
    except Exception:
        return info.filename


def _strip_archive_suffix(name):
    """去掉归档后缀，用于生成解压目录名。"""
    lower = name.lower()
    for suffix in STRIP_SUFFIXES:
        if lower.endswith(suffix):
            return name[: -len(suffix)]
    return name


def _unique_dir(path):
    """解压目录重名时依次追加 -2、-3，避免覆盖上一次解压结果。"""
    if not os.path.exists(path) or not os.listdir(path):
        return path
    index = 2
    while True:
        candidate = f"{path}-{index}"
        if not os.path.exists(candidate) or not os.listdir(candidate):
            return candidate
        index += 1


def _unpack_zip(path, dest_dir):
    """解压 zip，返回解压出的文件数。"""
    count = 0
    with zipfile.ZipFile(path) as zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            target = os.path.join(dest_dir, _zip_member_name(info).replace("\\", "/"))
            if not _is_within(dest_dir, target):
                continue
            os.makedirs(os.path.dirname(target) or dest_dir, exist_ok=True)
            with zf.open(info) as src, open(target, "wb") as dst:
                shutil.copyfileobj(src, dst)
            count += 1
    return count


def _unpack_tar(path, dest_dir):
    """解压 tar / tar.gz / tgz 等，返回解压出的文件数。"""
    count = 0
    with tarfile.open(path) as tf:
        for member in tf.getmembers():
            if not member.isfile():
                continue
            target = os.path.join(dest_dir, member.name)
            if not _is_within(dest_dir, target):
                continue
            source = tf.extractfile(member)
            if source is None:
                continue
            os.makedirs(os.path.dirname(target) or dest_dir, exist_ok=True)
            with source, open(target, "wb") as dst:
                shutil.copyfileobj(source, dst)
            count += 1
    return count


def _unpack_single(path, dest_dir):
    """解压单文件压缩（.gz/.bz2/.xz），返回解压出的文件数。"""
    suffix = os.path.splitext(path)[1].lower()
    module = SINGLE_COMPRESS_MODULES[suffix]
    name = os.path.basename(path)[: -len(suffix)] or "unpacked"
    target = os.path.join(dest_dir, name)
    with module.open(path, "rb") as src, open(target, "wb") as dst:
        shutil.copyfileobj(src, dst)
    return 1


def _unpack_external(kind, path, dest_dir):
    """用系统命令解压 rar / 7z，返回 (是否成功, 说明)。"""
    errors = []
    for label, template in EXTERNAL_UNPACKERS.get(kind, ()):
        executable = shutil.which(label)
        if not executable:
            errors.append(f"{label}=未安装")
            continue
        argv = [executable] + [
            part.format(archive=os.path.abspath(path), dest=os.path.abspath(dest_dir))
            for part in template
        ]
        try:
            proc = subprocess.run(
                argv, capture_output=True, text=True,
                encoding="utf-8", errors="replace", timeout=300,
            )
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{label}=执行失败({exc})")
            continue
        if proc.returncode == 0:
            return True, label
        detail = (proc.stderr or proc.stdout or "").strip().replace("\n", " ")[:200]
        errors.append(f"{label}=退出码 {proc.returncode}: {detail}")
    return False, "；".join(errors) or "没有可用的解压命令"


def unpack_archive(path, dest_dir, kind):
    """解压归档到 dest_dir，返回 (是否成功, 说明文本)。"""
    os.makedirs(dest_dir, exist_ok=True)
    try:
        if kind == "zip":
            return True, f"已解压 {_unpack_zip(path, dest_dir)} 个文件"
        if kind == "tar":
            return True, f"已解压 {_unpack_tar(path, dest_dir)} 个文件"
        if kind == "single":
            return True, f"已解压 {_unpack_single(path, dest_dir)} 个文件"
        ok, detail = _unpack_external(kind, path, dest_dir)
        if not ok:
            return False, detail
        count = sum(len(files) for _root, _dirs, files in os.walk(dest_dir))
        return True, f"已用 {detail} 解压 {count} 个文件"
    except Exception as exc:  # noqa: BLE001
        return False, f"{type(exc).__name__}: {exc}"


def extract_one(path, max_chars, display_name=None):
    """抽取单个文件的文本，返回 (标题, 正文)。display_name 用于标明文件在压缩包内的位置。"""
    suffix = os.path.splitext(path)[1].lower()
    name = display_name or os.path.basename(path)
    size = os.path.getsize(path) if os.path.exists(path) else 0
    header = f"{name} | {suffix or '无扩展名'} | {size} B"
    try:
        if suffix in TEXT_SUFFIXES or suffix == "":
            body = read_text_file(path)
        elif suffix == ".docx":
            body = read_docx(path)
        elif suffix == ".pdf":
            body = read_pdf(path)
        elif suffix in {".xlsx", ".xlsm"}:
            body = read_xlsx(path)
        elif suffix == ".pptx":
            body = read_pptx(path)
        elif suffix in IMAGE_SUFFIXES:
            body = "[图片附件] 请使用图像查看能力查看该文件，不要臆测内容：" + path
        elif suffix in UNSUPPORTED_HINT:
            body = "[无法解析] " + UNSUPPORTED_HINT[suffix]
        else:
            body = f"[无法解析] 不支持的附件格式 {suffix}，请人工查看或转换格式"
    except Exception as exc:  # noqa: BLE001
        body = f"[解析异常] {type(exc).__name__}: {exc}"
    body = body.strip()
    if max_chars and len(body) > max_chars:
        body = body[:max_chars] + f"\n…（已截断，全文共 {len(body)} 字）"
    return header, body


def _handle_entry(path, unpack_root, notes, depth, display_prefix):
    """处理单个条目：普通文件直接返回，压缩包先解压再递归解析。"""
    name = os.path.basename(path)
    kind = archive_kind(path)
    if kind is None:
        return [(path, display_prefix + name)]

    size = os.path.getsize(path) if os.path.exists(path) else 0
    head = f"{display_prefix}{name} | {os.path.splitext(name)[1].lower() or '无扩展名'} | {size} B"
    if depth >= MAX_ARCHIVE_DEPTH:
        notes.append(f"===== 附件：{head} =====\n[压缩包] 嵌套超过 {MAX_ARCHIVE_DEPTH} 层，已停止解压，请人工查看")
        return []

    prefix_part = "_".join(
        _strip_archive_suffix(seg) for seg in display_prefix.split("/") if seg
    )
    dir_name = (prefix_part + "_" if prefix_part else "") + _strip_archive_suffix(name)
    dest_dir = _unique_dir(os.path.join(unpack_root, dir_name))
    ok, detail = unpack_archive(path, dest_dir, kind)
    if not ok:
        notes.append(
            f"===== 附件：{head} =====\n[压缩包解压失败] {detail}"
            "；请人工下载查看该附件，不要当成“附件没有内容”"
        )
        return []
    notes.append(
        f"===== 附件：{head} =====\n[压缩包] {detail}（解压到 {dest_dir}），以下为解压后各文件的解析结果"
    )
    return collect_files(dest_dir, unpack_root, notes, depth + 1, display_prefix + name + "/")


def collect_files(target, unpack_root, notes, depth=0, display_prefix=""):
    """展开目标为 [(真实路径, 展示名)]，压缩包自动解压后递归解析。"""
    results = []
    if os.path.isdir(target):
        excluded = os.path.abspath(unpack_root)
        for root, dirs, files in os.walk(target):
            dirs[:] = sorted(
                name for name in dirs
                if name != ARCHIVE_DIR_NAME
                and os.path.abspath(os.path.join(root, name)) != excluded
            )
            for name in sorted(files):
                results.extend(
                    _handle_entry(os.path.join(root, name), unpack_root, notes, depth, display_prefix)
                )
    else:
        results.extend(_handle_entry(target, unpack_root, notes, depth, display_prefix))
    return results


def force_utf8_stdout():
    """Windows 控制台默认 GBK，遇到 emoji 或生僻字会导致 print 崩溃，这里统一改成 UTF-8。"""
    for stream in ("stdout", "stderr"):
        target = getattr(sys, stream, None)
        if target is None:
            continue
        try:
            target.reconfigure(encoding="utf-8", errors="replace")
        except Exception:  # noqa: BLE001
            pass


def main():
    force_utf8_stdout()
    parser = argparse.ArgumentParser(description="抽取 Jira 附件文本")
    parser.add_argument("target", help="附件文件或目录")
    parser.add_argument("--max-chars", type=int, default=20000, help="单个附件最大输出字符数")
    parser.add_argument("--out", help="同时写入指定文件")
    parser.add_argument("--unpack-dir", help="压缩包解压目录，默认放在附件目录下的 __unpacked 子目录")
    args = parser.parse_args()

    if os.path.isdir(args.target):
        default_root = os.path.join(os.path.abspath(args.target), ARCHIVE_DIR_NAME)
    else:
        default_root = os.path.join(os.path.dirname(os.path.abspath(args.target)), ARCHIVE_DIR_NAME)
    unpack_root = os.path.abspath(args.unpack_dir) if args.unpack_dir else default_root

    notes = []
    pairs = collect_files(args.target, unpack_root, notes)
    chunks = list(notes)
    for path, display_name in pairs:
        header, body = extract_one(path, args.max_chars, display_name)
        chunks.append(f"===== 附件：{header} =====\n{body or '（内容为空）'}")

    output = "\n\n".join(chunks) if chunks else "[未找到任何文件]"
    print(output)
    if args.out:
        with open(args.out, "w", encoding="utf-8", newline="\n") as fh:
            fh.write(output)
    return 0


if __name__ == "__main__":
    sys.exit(main())
